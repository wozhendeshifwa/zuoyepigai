"""生成复核表：汇总页（逐题得分 + 区间 + 调分后）与逐题页（把握度 + 手写转录）。"""
import glob
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import calibrate
import config

HEAD_FILL = PatternFill('solid', fgColor='EEF2F7')


def names_from_pdfs(sub_dir):
    out = {}
    for p in glob.glob(os.path.join(sub_dir, '*.pdf')):
        base = os.path.splitext(os.path.basename(p))[0]
        if '_' in base:
            no, name = base.split('_', 1)
            out[no] = name
    return out


def load_grades(grade_dir):
    rows = []
    for p in sorted(glob.glob(os.path.join(grade_dir, '*.json'))):
        if os.path.basename(p).startswith('_'):
            continue  # 备份/中间文件不算一份判分
        d = json.load(open(p, encoding='utf-8'))
        rows.append({
            'stuNo': d['stuNo'],
            'raw': float(d['score']),
            'per': {str(i.get('no')): i for i in d.get('items', [])},
            'tokens': (d.get('usage') or {}).get('total_tokens', 0),
        })
    return sorted(rows, key=lambda r: r['stuNo'])


def build(grade_dir, sub_dir, out_xlsx, target=None, floor=None, cap=None,
          questions=None):
    questions = questions or config.QUESTIONS
    target = config.TARGET_MEAN if target is None else target
    floor = config.SCORE_FLOOR if floor is None else floor
    cap = config.SCORE_CAP if cap is None else cap

    names = names_from_pdfs(sub_dir)
    rows = load_grades(grade_dir)
    if not rows:
        raise SystemExit(f'{grade_dir} 里没有判分结果')
    for r in rows:
        r['name'] = names.get(r['stuNo'], '')
        r['total'] = calibrate.to_range(r['raw'], floor, cap)

    shift, adjusted = calibrate.calibrate([r['total'] for r in rows],
                                          target, floor, cap)
    mean_raw = sum(r['total'] for r in rows) / len(rows)
    mean_adj = sum(adjusted) / len(adjusted)

    wb = Workbook()
    ws = wb.active
    ws.title = '汇总'
    head = ['学号', '姓名'] + [f'{no}({full})' for no, _, full in questions] + \
           ['百分制小计', f'计入{floor:g}-{cap:g}', '调分后', '低把握题数']
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(1, c).font = Font(bold=True)
        ws.cell(1, c).fill = HEAD_FILL

    for r, adj in zip(rows, adjusted):
        line = [r['stuNo'], r['name']]
        for no, _, _ in questions:
            it = r['per'].get(no)
            line.append(float(it.get('score')) if it else 0.0)
        low = sum(1 for no, _, _ in questions
                  if (r['per'].get(no) or {}).get('conf') == '低')
        line += [round(r['raw'], 2), round(r['total'], 2), round(adj, 2), low]
        ws.append(line)

    ws.append([])
    ws.append(['均分', round(mean_raw, 2), '平移量', round(shift, 2),
               '调分后均分', round(mean_adj, 2),
               '目标', f'{target:g}±1', '区间', f'{floor:g}-{cap:g}'])
    for c in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.freeze_panes = 'C2'

    ws2 = wb.create_sheet('逐题')
    ws2.append(['学号', '姓名', '题号', '题型', '满分', 'AI 给分', '把握度',
                '学生作答（转录）'])
    for c in range(1, 9):
        ws2.cell(1, c).font = Font(bold=True)
        ws2.cell(1, c).fill = HEAD_FILL
    for r in rows:
        for no, qtype, full in questions:
            it = r['per'].get(no) or {}
            ws2.append([r['stuNo'], r['name'], no, qtype, full,
                        float(it.get('score') or 0), it.get('conf', ''),
                        it.get('transcript', '')])
    for c, w in zip(range(1, 9), [14, 10, 6, 8, 6, 8, 8, 90]):
        ws2.column_dimensions[get_column_letter(c)].width = w
    for row in ws2.iter_rows(min_row=2, min_col=8, max_col=8):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws2.freeze_panes = 'C2'

    os.makedirs(os.path.dirname(os.path.abspath(out_xlsx)), exist_ok=True)
    wb.save(out_xlsx)

    final = [{'stuNo': r['stuNo'], 'name': r['name'], 'raw': r['raw'],
              'total': r['total'], 'final': adj}
             for r, adj in zip(rows, adjusted)]
    with open(os.path.join(grade_dir, '_final.json'), 'w', encoding='utf-8') as f:
        json.dump(final, f, ensure_ascii=False, indent=1)

    return {
        'rows': rows, 'adjusted': adjusted, 'shift': shift,
        'mean_raw': mean_raw, 'mean_adj': mean_adj,
        'tokens': sum(r['tokens'] for r in rows), 'out': out_xlsx,
    }
